# -*- coding: utf-8 -*-

from odoo import fields, api, models, _
from odoo.exceptions import ValidationError
from datetime import date, datetime, time, timedelta
from calendar import monthrange
from openpyxl import load_workbook
from openpyxl.styles import Font, borders, Alignment, PatternFill
import base64
from io import BytesIO
from dateutil.relativedelta import relativedelta
from operator import itemgetter
from pytz import timezone

from .mailmerge import MailMerge
import base64
from io import BytesIO
from odoo.tools import float_is_zero, float_compare, pycompat

import logging

class SHDepositDocument(models.TransientModel):
    _name = 'sh.deposit.document'
    _description = 'Phiếu hẹn làm dịch vụ'

    services = fields.Text('Dịch vụ')
    attachs = fields.Text('Kèm theo - Hẹn lịch')
    counselor = fields.Many2one('res.users', string='Nhân viên hướng dẫn')
    note = fields.Text('Ghi chú')

    payment = fields.Many2one('account.payment', 'Phiếu thanh toán')

    # ACTION XUAT PHIEU HEN
    def action_report_phieu_hen(self):
        surgery_attacment = self.env.ref('shealth_all_in_one.phieuhen_report_attachment')
        decode = base64.b64decode(surgery_attacment.datas)
        doc = MailMerge(BytesIO(decode))
        # gia tri tra ve
        data_list = []

        record = self.payment

        record_data = {}
        record_data['TEN_CTY'] = str(record.company_id.name).upper()

        dia_chi = ''
        if record.company_id.street:
            dia_chi += record.company_id.street
        if record.company_id.street2:
            dia_chi += ', ' + record.company_id.street2
        if record.company_id.state_id:
            dia_chi += ', ' + record.company_id.state_id.name
        if record.company_id.country_id:
            dia_chi += ', ' + record.company_id.country_id.name
        record_data['DIA_CHI'] = dia_chi

        record_data['NGAY_THANH_TOAN'] = record.payment_date.strftime(
            '%d/%m/%Y') if record.payment_date else ''

        record_data['MA_KHACH_HANG'] = record.partner_id.code_customer
        record_data['MA_BOOKING'] = record.crm_id.name
        record_data['TEN_KHACH_HANG'] = record.partner_id.name

        dia_chi_kh = ''
        if record.partner_id.street:
            dia_chi_kh += record.partner_id.street
        if record.partner_id.district_id:
            dia_chi_kh += ', ' + record.partner_id.district_id.name
        if record.partner_id.state_id:
            dia_chi_kh += ', ' + record.partner_id.state_id.name
        if record.partner_id.country_id:
            dia_chi_kh += ', ' + record.partner_id.country_id.name
        record_data['DIA_CHI_KH'] = dia_chi_kh

        # LẦN THANH TOÁN THEO BOOKING VÀ THEO NGÀY
        lan_thanh_toan = ''
        tong_nhan = 0.00
        # payments = record.crm_id.payment_ids.filtered(lambda p: p.payment_date == '2021-11-19')
        payments = record.crm_id.payment_ids.filtered(lambda p: p.state in ['posted', 'reconciled'] and p.payment_type in ['inbound','outbound'] and p.payment_date == record.payment_date)

        if len(payments) > 0:
            for pay in payments:
                payment_method = str(dict(pay._fields['payment_method']._description_selection(self.env)).get(
                    pay.payment_method))
                payment_type = str(dict(pay._fields['payment_type']._description_selection(self.env)).get(
                    pay.payment_type))

                lan_thanh_toan += "* " + str(pay.payment_date.strftime(
                    '%d/%m/%Y') if pay.payment_date else '') + ": (" + payment_type + " - " + str(
                    payment_method) + ") " + str(
                    "{:,.2f}".format(pay.amount)) + " " + str(pay.currency_id.name)

                if pay.currency_rate_id:
                    rate = str("{:,.2f}".format(round(1/pay.currency_rate_id.rate)))
                    lan_thanh_toan += " - Tỷ giá: " + rate + " VNĐ \n"
                else:
                    lan_thanh_toan += "\n"

                if pay.payment_type == 'inbound':
                    tong_nhan += pay.amount_vnd
                else:
                    tong_nhan -= pay.amount_vnd

        record_data['LAN_THANH_TOAN'] = lan_thanh_toan + '\n Tổng nhận: ' + str("{:,.0f}".format(tong_nhan)) + 'VNĐ \n'
        record_data['DICH_VU'] = self.services
        record_data['KEM_THEO'] = self.attachs
        record_data['NV_TU_VAN'] = self.counselor.name
        record_data['GHI_CHU'] = self.note
        record_data['NGUOI_THU'] = self.env.user.name

        data_list.append(record_data)
        doc.merge_templates(data_list, separator='page_break')

        fp = BytesIO()
        doc.write(fp)
        doc.close()
        fp.seek(0)
        report = base64.encodebytes((fp.read()))
        fp.close()
        attachment = self.env['ir.attachment'].sudo().create({'name': 'PHIEU_HEN.docx',
                                                              'datas': report,
                                                              'res_model': 'temp.creation',
                                                              'public': True})

        # return {'name': 'PHIẾU HẸN',
        #                     'type': 'ir.actions.act_window',
        #                     'res_model': 'temp.wizard',
        #                     'view_mode': 'form',
        #                     'target': 'inline',
        #                     'view_id': self.env.ref('ms_templates.report_wizard').id,
        #                     'context': {'attachment_id': attachment.id}
        #             }

        url = "/web/content/?model=ir.attachment&id=%s&filename_field=name&field=datas&download=true" \
              % attachment.id
        return {'name': 'PHIẾU HẸN',
                'type': 'ir.actions.act_url',
                'url': url,
                'target': 'self',
                }